# ingestion/docx_chargeur.py
from pathlib import Path
from typing import List
from docx import Document as DocxDocument
from models.document import Document

def extraire_texte_docx(chemin_fichier: Path) -> str:
    """Lit un .docx et renvoie tout le texte brut (paragraphes concaténés)."""
    docx_obj = DocxDocument(chemin_fichier)
    paragraphes = [p.text.strip() for p in docx_obj.paragraphs if p.text.strip()]
    return "\n".join(paragraphes)

def charger_document_docx(chemin_fichier: str) -> Document:
    """Crée un objet Document à partir d’un fichier .docx."""
    chemin = Path(chemin_fichier)
    texte = extraire_texte_docx(chemin)
    meta = {
        "type": "docx",
        "nom_fichier": chemin.name,
    }
    return Document(texte=texte, metadonnees=meta)
